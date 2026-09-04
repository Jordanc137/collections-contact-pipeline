# =============================================================================
# step1_consolidate.py - CDR Ingestion (Incremental)
#
# Reads raw call detail records from all configured input folders, normalizes
# them to a single schema, and appends new records to cdr_master.parquet.
#
# SUPPORTED INPUT FORMATS (detected by column headers — NEVER by filename):
#   A) Systems XLSX  : contains "DURACION SEG" + "CATEGORIA LLAMADA"
#   B) Operations XLSX : "FECHA INICIO Y HORA" | "PHONE NUMBER" | "STATUS" | "CAMPAIGN ID" | "CUENTA"
#   C) Vicidial TXT  : tab-delimited, "call_date" + "phone_number_dialed"
#   D) AI calls CSV  : "campaign_date" + "answered_by" + "contact_f_id"
#   E) WhatsApp CSV  : "Teléfono" + "Resultado" + "Fecha de inicio"
#   F) Blaster CHOCK CSV : "Estado" + "Duración" + "telefono" + "Campaign ID"
#   G) MuttechMX XLSX: "FechaLlamada" + "TelefonoMarcado" + "Estatus"
#
# INCREMENTAL DEDUPLICATION:
# Keyed on (DATE, TOOL) pairs — NOT just DATE. Using only DATE as the key
# silently drops entire tools when files for different tools arrive on
# different days for the same campaign date. The composite key is loaded
# from the existing Parquet using columnar reads (3 columns, not the full file).
#
# OUTPUT SCHEMA (cdr_master.parquet):
#   PHONE, DATE, HOUR, DATETIME, DURATION_SEC, HAS_DURATION, STATUS,
#   CALL_CATEGORY, TOOL, CAMPAIGN_ID, CAMPAIGN_ID_CDR, ACCOUNT_ID,
#   IS_RPC, IS_MESSAGE, SOURCE_FILE
# =============================================================================
import csv
import os
import re
import sys
import glob
import time
from collections import defaultdict

# Add project root to path for config and core imports
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config
from core import cdr_io

try:
    from python_calamine import CalamineWorkbook
    _HAS_CALAMINE = True
except ImportError:
    _HAS_CALAMINE = False

# ---------------------------------------------------------------------------
# OUTPUT SCHEMA
# ---------------------------------------------------------------------------
FIELDNAMES = [
    "PHONE", "DATE", "HOUR", "DATETIME", "DURATION_SEC", "HAS_DURATION",
    "STATUS", "CALL_CATEGORY", "TOOL", "CAMPAIGN_ID", "CAMPAIGN_ID_CDR",
    "ACCOUNT_ID", "IS_RPC", "IS_MESSAGE", "SOURCE_FILE",
]

MASTER_PATH = os.path.join(config.OUTPUT_FOLDER, config.MASTER_PARQUET)

# ---------------------------------------------------------------------------
# STATUS CATALOGS
# Extend these for your dialing platforms. Unknown statuses fall to "No data"
# and trigger a console warning — never silently miscategorized.
# ---------------------------------------------------------------------------

# Predictive dialer (Vicidial-based)
CAT_PREDICTIVE = {
    # Answered — agent confirmed contact
    "COL": "Answered", "CTT": "Answered", "SALE": "Answered",
    "PDPP": "Answered", "CALLBK": "Answered", "XFER": "Answered",
    "NA":   "Answered",  # as of 2026-09-02 confirmed = answered
    # Ringing / busy / dropped
    "AB": "Ringing", "ADCA": "Ringing", "ADC": "Ringing",
    "DROP": "Ringing", "PDROP": "Ringing", "SVYCLM": "Ringing",
    "ILO": "Ringing", "MSB": "Ringing", "NOA": "Ringing",
    "REF": "Ringing", "TIT": "Ringing", "MSJ": "Ringing",
    "FAM": "Ringing",
    # Discard — disconnected, answering machine, invalid
    "AA": "Discard", "SINLIN": "Discard", "DEF": "Discard",
    "BZNDIR": "Discard", "FVN": "Discard",
}

# Robocall / Blaster (Vicidial-based, same platform)
CAT_BLASTER = dict(CAT_PREDICTIVE)

# Alternate robocall platform (Blaster CHOCK)
CAT_BLASTER_CHOCK = {
    "CONTESTADA": "Answered", "COMPLETADO": "Answered", "TRANSFERIDO": "Answered",
    "NO CONTESTA": "Ringing", "BUZÓN": "Ringing", "BUZON": "Ringing",
    "RECHAZADA": "Ringing",
}
MESSAGE_STATUSES_BLASTER_CHOCK = {"COMPLETADO", "TRANSFERIDO"}

# Alternate predictive platform (MuttechMX)
CAT_MUTTECHMX = {
    "AB": "Ringing", "AA": "Discard", "OK": "Answered",
    "ADC": "Ringing", "DROP": "Ringing", "PDROP": "Ringing",
    "ERI": "No data",  # "Agent Error" — pending business confirmation
    "NA":  "No data",  # "LLAMADA CONTESTADA" but 0 duration + system user — ambiguous
}

# RPC (Right Party Contact) statuses — Predictive only
RPC_STATUSES = {"COL", "CTT", "SALE", "PDPP", "CALLBK", "XFER", "NA"}
# IS_MESSAGE proxy statuses — Blaster/IVR (machine delivered a message)
MESSAGE_STATUSES = {"XFER", "PM", "PU", "SVYCLM"}

# ---------------------------------------------------------------------------
# XLSX reader wrapper (calamine → openpyxl fallback)
# ---------------------------------------------------------------------------
class _XlsxReader:
    """Opens an XLSX file once and exposes sheet_names + a rows() iterator.
    Supports both python-calamine (fast, Rust-based) and openpyxl (fallback).
    The same open file object is used for both header detection and data
    iteration — no double-open.
    """
    def __init__(self, path):
        self.path = path
        self._wb_openpyxl = None
        if _HAS_CALAMINE:
            try:
                self._wb = CalamineWorkbook.from_path(path)
                self.sheet_names = list(self._wb.sheet_names)
                self._mode = "calamine"
                return
            except Exception:
                pass
        import openpyxl
        self._wb_openpyxl = openpyxl.load_workbook(path, read_only=True)
        self._wb = self._wb_openpyxl
        self.sheet_names = list(self._wb.sheetnames)
        self._mode = "openpyxl"

    def rows(self, sheet_name):
        if self._mode == "calamine":
            return iter(self._wb.get_sheet_by_name(sheet_name).iter_rows())
        return self._wb[sheet_name].iter_rows(values_only=True)

    def close(self):
        if self._wb_openpyxl:
            self._wb_openpyxl.close()


class _DelimReader:
    """Generic reader for tab/comma-delimited text files."""
    def __init__(self, path, delimiter=","):
        self.path = path
        self.sheet_names = [None]
        self.delimiter = delimiter
        self._f = open(path, encoding="utf-8-sig", newline="")

    def rows(self, _sheet=None):
        return csv.reader(self._f, delimiter=self.delimiter)

    def close(self):
        self._f.close()


# ---------------------------------------------------------------------------
# Phone normalization
# ---------------------------------------------------------------------------
_RE_NON_DIGIT = re.compile(r"\D")

def norm_phone(raw):
    if raw is None:
        return None
    s = str(raw).strip().split(".")[0]
    s = _RE_NON_DIGIT.sub("", s)
    if len(s) < 10:
        return None
    return s[-10:]


# ---------------------------------------------------------------------------
# Tool identification from campaign ID
# ---------------------------------------------------------------------------
def tool_from_campaign_id(cid):
    c = (cid or "").strip().upper()
    for tool, prefixes in config.CAMPAIGN_ID_PREFIXES.items():
        if any(c.startswith(p.upper()) for p in prefixes):
            return tool
    return "Unknown"


def _is_blaster_chock_file(filename):
    """Blaster CHOCK files are .csv with a specific naming pattern.
    Used only for dedup key disambiguation — not for format detection."""
    low = filename.lower()
    return low.endswith(".csv") and "tel" in low

def _is_muttechmx_file(filename):
    low = filename.lower()
    return low.endswith(".xlsx") and "reporte" in low


# ---------------------------------------------------------------------------
# Dedup key
# ---------------------------------------------------------------------------
def _dedup_key(date, tool, source_file):
    """Composite dedup key: (DATE, TOOL_SOURCE).
    Blaster CHOCK and MuttechMX share a business TOOL name with their classic
    counterparts but are independent sources — disambiguate for dedup.
    """
    t = tool
    if t == "Blaster" and _is_blaster_chock_file(source_file):
        t = "Blaster_CHOCK"
    elif t == "Predictivo" and _is_muttechmx_file(source_file):
        t = "Predictivo_MT"
    return (date, t)


def load_existing_keys():
    """Load (DATE, TOOL) pairs already in the master using columnar read."""
    keys = set()
    if os.path.exists(MASTER_PATH):
        cols = cdr_io.read_parquet_columns(
            MASTER_PATH, ["DATE", "TOOL", "SOURCE_FILE"]
        )
        for d, t, o in zip(cols["DATE"], cols["TOOL"], cols["SOURCE_FILE"]):
            keys.add(_dedup_key(d, t, o or ""))
    return keys


# ---------------------------------------------------------------------------
# Format parsers — one per input format
# Each returns the number of rows written.
# ---------------------------------------------------------------------------

def parse_systems_xlsx(rows_iter, path, writer, existing_keys, stats):
    """Format A: Systems XLSX with pre-computed CALL_CATEGORY."""
    origin = os.path.basename(path)
    idx = None
    n = 0
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        if idx is None:
            idx = {h: i for i, h in enumerate(row)}
            continue
        date = row[idx.get("FECHA", -1)][:10] if "FECHA" in idx else ""
        phone = norm_phone(row[idx.get("PHONE NUMBER", -1)])
        status = row[idx.get("STATUS", -1)]
        cid = row[idx.get("CAMPAIGN ID", -1)]
        tool = tool_from_campaign_id(cid)
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        cat = row[idx.get("CATEGORIA LLAMADA", -1)] or "No data"
        dur_raw = row[idx.get("DURACION SEG", -1)]
        try:
            dur = int(float(dur_raw)) if dur_raw else None
        except ValueError:
            dur = None
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": "", "DATETIME": "",
            "DURATION_SEC": dur if dur is not None else "",
            "HAS_DURATION": 1 if dur is not None else 0,
            "STATUS": status, "CALL_CATEGORY": cat,
            "TOOL": tool, "CAMPAIGN_ID": cid,
            "CAMPAIGN_ID_CDR": cid, "ACCOUNT_ID": "",
            "IS_RPC": 1 if status in RPC_STATUSES else 0,
            "IS_MESSAGE": 1 if status in MESSAGE_STATUSES else 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    return n


def parse_operations_xlsx(rows_iter, header, path, writer, existing_keys, stats):
    """Format B: Operations XLSX — no duration, category derived from STATUS catalog."""
    origin = os.path.basename(path)
    idx = {h: i for i, h in enumerate(header)}
    required = ("FECHA INICIO Y HORA", "PHONE NUMBER", "STATUS", "CAMPAIGN ID", "CUENTA")
    if any(c not in idx for c in required):
        print(f"  [SKIP] {origin}: missing required columns for format B")
        return 0
    n = 0
    seen_in_run = set()
    unknown_statuses = set()
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        datetime_raw = row[idx["FECHA INICIO Y HORA"]]
        date = datetime_raw[:10] if datetime_raw else ""
        hour = datetime_raw[11:13] if len(datetime_raw) > 13 else ""
        phone = norm_phone(row[idx["PHONE NUMBER"]])
        status = row[idx["STATUS"]].upper()
        cid = row[idx["CAMPAIGN ID"]]
        account = row[idx["CUENTA"]]
        tool = tool_from_campaign_id(cid)
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        # Row-level dedup within the same file (same datetime+phone+cid)
        row_key = (datetime_raw, phone, cid)
        if row_key in seen_in_run:
            stats["skipped_row_dedup"] += 1
            continue
        seen_in_run.add(row_key)
        cat = CAT_PREDICTIVE.get(status)
        if cat is None and status:
            unknown_statuses.add(status)
            cat = "No data"
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": hour, "DATETIME": datetime_raw,
            "DURATION_SEC": "", "HAS_DURATION": 0,
            "STATUS": status, "CALL_CATEGORY": cat or "No data",
            "TOOL": tool, "CAMPAIGN_ID": cid,
            "CAMPAIGN_ID_CDR": cid, "ACCOUNT_ID": account,
            "IS_RPC": 1 if status in RPC_STATUSES else 0,
            "IS_MESSAGE": 1 if status in MESSAGE_STATUSES else 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    if unknown_statuses:
        print(f"  [WARNING] {origin}: unknown STATUS codes → 'No data': {sorted(unknown_statuses)}")
    return n


def parse_vicidial_txt(rows_iter, header, path, writer, existing_keys, stats):
    """Format C: Vicidial EXPORT_CALL_REPORT (tab-delimited)."""
    origin = os.path.basename(path)
    idx = {h: i for i, h in enumerate(header)}
    n = 0
    unknown_statuses = set()
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        date = row[idx.get("call_date", -1)][:10] if "call_date" in idx else ""
        phone = norm_phone(row[idx.get("phone_number_dialed", "")])
        status = row[idx.get("status", "")].upper()
        cid = row[idx.get("campaign_id", "")]
        tool = tool_from_campaign_id(cid)
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        cat = CAT_PREDICTIVE.get(status)
        if cat is None and status:
            unknown_statuses.add(status)
            cat = "No data"
        dur_raw = row[idx.get("length_in_sec", "")]
        try:
            dur = int(float(dur_raw)) if dur_raw else None
        except ValueError:
            dur = None
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": row[idx.get("call_time", "")][:2],
            "DATETIME": f"{date} {row[idx.get('call_time','')]}".strip(),
            "DURATION_SEC": dur if dur is not None else "",
            "HAS_DURATION": 1 if dur is not None else 0,
            "STATUS": status, "CALL_CATEGORY": cat or "No data",
            "TOOL": tool, "CAMPAIGN_ID": cid,
            "CAMPAIGN_ID_CDR": cid,
            "ACCOUNT_ID": row[idx.get("lead_id", "")],
            "IS_RPC": 1 if status in RPC_STATUSES else 0,
            "IS_MESSAGE": 1 if status in MESSAGE_STATUSES else 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    if unknown_statuses:
        print(f"  [WARNING] {origin}: unknown STATUS codes → 'No data': {sorted(unknown_statuses)}")
    return n


def parse_ai_csv(rows_iter, header, path, writer, existing_keys, stats):
    """Format D: AI/chat calls (calls_multiple_campaigns*.csv)."""
    origin = os.path.basename(path)
    idx = {h: i for i, h in enumerate(header)}
    n = 0
    tool = "IA"
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        date = row[idx.get("campaign_date", "")][:10]
        phone = norm_phone(row[idx.get("number_to", "")])
        status = row[idx.get("status", "")].upper()
        account = row[idx.get("contact_f_id", "")]
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        dur_raw = row[idx.get("duration", "")]
        try:
            dur = int(float(dur_raw)) if dur_raw else None
        except ValueError:
            dur = None
        answered = status == "ANSWERED"
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": row[idx.get("started_at", "")][11:13],
            "DATETIME": row[idx.get("started_at", "")],
            "DURATION_SEC": dur if dur is not None else "",
            "HAS_DURATION": 1 if dur is not None else 0,
            "STATUS": status,
            "CALL_CATEGORY": "Answered" if answered else "Ringing",
            "TOOL": tool, "CAMPAIGN_ID": row[idx.get("campaign_name", "")],
            "CAMPAIGN_ID_CDR": "", "ACCOUNT_ID": account,
            "IS_RPC": 0, "IS_MESSAGE": 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    return n


def parse_whatsapp_csv(rows_iter, header, path, writer, existing_keys, stats):
    """Format E: WhatsApp / chat platform CSV."""
    origin = os.path.basename(path)
    idx = {h: i for i, h in enumerate(header)}
    n = 0
    tool = "WhatsApp"
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        date_raw = row[idx.get("Fecha de inicio", "")]
        date = date_raw[:10].replace("/", "-") if date_raw else ""
        phone = norm_phone(row[idx.get("Teléfono", "")])
        resultado = row[idx.get("Resultado", "")]
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        answered = resultado.upper() in ("RESUELTO", "CERRADO", "ATENDIDO", "RESOLVED", "CLOSED")
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": date_raw[11:13] if len(date_raw) > 13 else "",
            "DATETIME": date_raw,
            "DURATION_SEC": "", "HAS_DURATION": 0,
            "STATUS": resultado,
            "CALL_CATEGORY": "Answered" if answered else "Ringing",
            "TOOL": tool, "CAMPAIGN_ID": "",
            "CAMPAIGN_ID_CDR": "", "ACCOUNT_ID": "",
            "IS_RPC": 0, "IS_MESSAGE": 1 if answered else 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    return n


def parse_blaster_chock_csv(rows_iter, header, path, writer, existing_keys, stats):
    """Format F: Blaster CHOCK — alternate robocall platform."""
    origin = os.path.basename(path)
    idx = {h: i for i, h in enumerate(header)}
    tool = "Blaster"
    n = 0
    unknown_statuses = set()
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        date = row[idx.get("Fecha", "")][:10]
        hora = row[idx.get("Hora", "")]
        estado = row[idx.get("Estado", "")].upper()
        phone = norm_phone(row[idx.get("telefono", "")])
        cid = row[idx.get("Campaign ID", "")]
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        cat = CAT_BLASTER_CHOCK.get(estado)
        if cat is None and estado:
            unknown_statuses.add(estado)
            cat = "No data"
        dur_raw = row[idx.get("Duración", "")]
        try:
            dur = int(float(dur_raw)) if dur_raw else None
        except ValueError:
            dur = None
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": hora[:2] if hora else "",
            "DATETIME": f"{date} {hora}".strip(),
            "DURATION_SEC": dur if dur is not None else "",
            "HAS_DURATION": 1 if dur is not None else 0,
            "STATUS": estado, "CALL_CATEGORY": cat or "No data",
            "TOOL": tool, "CAMPAIGN_ID": cid,
            "CAMPAIGN_ID_CDR": "", "ACCOUNT_ID": "",
            "IS_RPC": 0,
            "IS_MESSAGE": 1 if estado in MESSAGE_STATUSES_BLASTER_CHOCK else 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    if unknown_statuses:
        print(f"  [WARNING] {origin}: unknown ESTADO codes → 'No data': {sorted(unknown_statuses)}")
    return n


def parse_muttechmx_xlsx(rows_iter, header, path, writer, existing_keys, stats):
    """Format G: MuttechMX — alternate predictive platform."""
    origin = os.path.basename(path)
    idx = {h: i for i, h in enumerate(header)}
    tool = "Predictivo"
    n = 0
    unknown_statuses = set()
    for raw_row in rows_iter:
        row = [str(c).strip() if c is not None else "" for c in raw_row]
        date_raw = row[idx.get("FechaLlamada", "")]
        date = str(date_raw)[:10]
        phone = norm_phone(row[idx.get("TelefonoMarcado", "")])
        status = row[idx.get("Estatus", "")].upper() if row[idx.get("Estatus", "")] else ""
        key = _dedup_key(date, tool, origin)
        if not date or key in existing_keys:
            stats["skipped_dedup"] += 1
            continue
        cat = CAT_MUTTECHMX.get(status)
        if cat is None and status:
            unknown_statuses.add(status)
            cat = "No data"
        dur_raw = row[idx.get("DuracionMinutos", "")]
        try:
            dur = int(float(dur_raw) * 60) if dur_raw else None
        except ValueError:
            dur = None
        writer.writerow({
            "PHONE": phone or "", "DATE": date,
            "HOUR": "", "DATETIME": date_raw,
            "DURATION_SEC": dur if dur is not None else "",
            "HAS_DURATION": 1 if dur is not None else 0,
            "STATUS": status, "CALL_CATEGORY": cat or "No data",
            "TOOL": tool, "CAMPAIGN_ID": "",
            "CAMPAIGN_ID_CDR": "", "ACCOUNT_ID": "",
            "IS_RPC": 1 if status == "OK" else 0,
            "IS_MESSAGE": 0,
            "SOURCE_FILE": origin,
        })
        n += 1
        stats[("tool", tool)] += 1
    if unknown_statuses:
        print(f"  [WARNING] {origin}: unknown Estatus codes → 'No data': {sorted(unknown_statuses)}")
    return n


# ---------------------------------------------------------------------------
# File detection
# ---------------------------------------------------------------------------
def detect_files():
    """Scan all input folders and return a list of
    (path, format_type, sheet_name, reader, row_iterator, header) tuples.
    Detection is always by column headers — never by filename.
    """
    found = []
    all_dirs = [
        os.path.join(config.PIPELINE_BASE, d) if d else config.PIPELINE_BASE
        for d in config.INPUT_FOLDERS
    ]

    # XLSX files
    xlsx_paths = []
    for d in all_dirs:
        if os.path.isdir(d):
            xlsx_paths += glob.glob(os.path.join(d, "*.xlsx"))
    for path in sorted(set(xlsx_paths)):
        if os.path.basename(path).startswith("~$"):
            continue
        try:
            reader = _XlsxReader(path)
        except Exception as e:
            print(f"  (skipped, could not open {os.path.basename(path)}: {e})")
            continue
        for sheet in reader.sheet_names:
            row_iter = reader.rows(sheet)
            try:
                first_row = next(row_iter)
            except StopIteration:
                continue
            header = [str(c).strip() if c is not None else "" for c in first_row]
            header_set = set(header)
            if "DURACION SEG" in header_set and "CATEGORIA LLAMADA" in header_set:
                found.append((path, "systems", sheet, reader, row_iter, header))
                break
            elif "FECHA INICIO Y HORA" in header_set and "PHONE NUMBER" in header_set:
                found.append((path, "operations", sheet, reader, row_iter, header))
                break
            elif "FechaLlamada" in header_set and "TelefonoMarcado" in header_set:
                found.append((path, "muttechmx", sheet, reader, row_iter, header))
                break
        else:
            reader.close()

    # TXT files (Vicidial EXPORT_CALL_REPORT, tab-delimited)
    txt_paths = []
    for d in all_dirs:
        if os.path.isdir(d):
            txt_paths += glob.glob(os.path.join(d, "*.txt"))
    for path in sorted(set(txt_paths)):
        try:
            reader = _DelimReader(path, delimiter="\t")
            row_iter = reader.rows()
            header = [c.strip() for c in next(row_iter)]
        except StopIteration:
            continue
        except Exception as e:
            print(f"  (skipped {os.path.basename(path)}: {e})")
            continue
        if "call_date" in header and "phone_number_dialed" in header:
            found.append((path, "vicidial_txt", None, reader, row_iter, header))
        else:
            reader.close()

    # CSV files (IA, WhatsApp, Blaster CHOCK)
    csv_paths = []
    for d in all_dirs:
        if os.path.isdir(d):
            csv_paths += glob.glob(os.path.join(d, "*.csv"))
    for path in sorted(set(csv_paths)):
        try:
            reader = _DelimReader(path, delimiter=",")
            row_iter = reader.rows()
            header = [c.strip() for c in next(row_iter)]
        except StopIteration:
            continue
        except Exception as e:
            print(f"  (skipped {os.path.basename(path)}: {e})")
            continue
        header_set = set(header)
        if "campaign_date" in header_set and "contact_f_id" in header_set:
            found.append((path, "ai_csv", None, reader, row_iter, header))
        elif "Teléfono" in header_set and "Resultado" in header_set and "Fecha de inicio" in header_set:
            found.append((path, "whatsapp_csv", None, reader, row_iter, header))
        elif "Estado" in header_set and "Duración" in header_set and "Campaign ID" in header_set:
            found.append((path, "blaster_chock", None, reader, row_iter, header))
        else:
            reader.close()

    return found


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    if not _HAS_CALAMINE:
        print("[INFO] python-calamine not installed — using openpyxl (slower). "
              "Run: pip install python-calamine")

    os.makedirs(config.OUTPUT_FOLDER, exist_ok=True)
    stats = defaultdict(int)
    existing_keys = load_existing_keys()

    if existing_keys:
        dates = sorted({d for d, _ in existing_keys})
        print(f"Master has data for {len(dates)} date(s): {dates[0]} … {dates[-1]}")

    files = detect_files()
    print(f"Input files detected: {len(files)}")

    mode = "a" if os.path.exists(MASTER_PATH) else "w"
    new_keys_this_run = set()

    def _track_key(row):
        new_keys_this_run.add(_dedup_key(row["DATE"], row["TOOL"], row["SOURCE_FILE"]))

    with cdr_io.ParquetDictWriter(MASTER_PATH, FIELDNAMES, mode=mode, on_row=_track_key) as w:
        w.writeheader()
        for path, fmt, sheet, reader, row_iter, header in files:
            t0 = time.time()
            if fmt == "systems":
                n = parse_systems_xlsx(row_iter, path, w, existing_keys, stats)
            elif fmt == "operations":
                n = parse_operations_xlsx(row_iter, header, path, w, existing_keys, stats)
            elif fmt == "vicidial_txt":
                n = parse_vicidial_txt(row_iter, header, path, w, existing_keys, stats)
            elif fmt == "ai_csv":
                n = parse_ai_csv(row_iter, header, path, w, existing_keys, stats)
            elif fmt == "whatsapp_csv":
                n = parse_whatsapp_csv(row_iter, header, path, w, existing_keys, stats)
            elif fmt == "blaster_chock":
                n = parse_blaster_chock_csv(row_iter, header, path, w, existing_keys, stats)
            elif fmt == "muttechmx":
                n = parse_muttechmx_xlsx(row_iter, header, path, w, existing_keys, stats)
            else:
                n = 0
            reader.close()
            if n > 0:
                existing_keys.update(new_keys_this_run)
            elapsed = time.time() - t0
            print(f"  [{fmt}] {os.path.basename(path)}: {n:,} rows ({elapsed:.1f}s)")

    print("\nRows by tool (this run):")
    for k, v in sorted(stats.items(), key=lambda x: str(x)):
        if isinstance(k, tuple) and k[0] == "tool":
            print(f"  {k[1]}: {v:,}")
    print(f"\nDone → {MASTER_PATH}")
